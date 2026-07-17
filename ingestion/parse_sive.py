"""
MA2E — Pilote Trésorerie
Parser du fichier "POINT SIVE" (suivi des prélèvements & reversements).

Expose parse_sive(path) -> dict (réutilisé par l'import live de l'API),
et un main() qui écrit le JSON pour le frontend (mode statique).
"""
from __future__ import annotations
import os, sys, json, datetime as dt
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "..", "POINT SIVE.xlsx")
OUT_DIR = os.path.join(HERE, "..", "frontend", "src", "data")
OUT = os.path.join(OUT_DIR, "reversements.json")

MOIS_FR = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
           "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]


def clean(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace(",", "")
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def periode_label(v):
    if isinstance(v, dt.datetime):
        return f"{MOIS_FR[v.month]} {v.year}", v.date().isoformat()
    return str(v).strip(), None


def parse_sive(path):
    """Parse un classeur POINT SIVE et renvoie le dict complet."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    lignes = []
    for row in rows:
        if not row or row[0] is None:
            continue
        c0 = row[0]
        is_date = isinstance(c0, dt.datetime)
        is_solde = isinstance(c0, str) and "solde" in c0.lower()
        if not (is_date or is_solde):
            continue
        total_fichier = clean(row[7]) if len(row) > 7 else None
        montant_vir = clean(row[8]) if len(row) > 8 else None
        if total_fichier is None and montant_vir is None:
            continue
        label, iso = periode_label(c0)
        ecart = (total_fichier or 0) - (montant_vir or 0)
        obs = (str(row[10]).strip() if len(row) > 10 and row[10] else
               ("OK" if abs(ecart) < 1 else "En attente de reversement"))
        lignes.append({
            "periode": label, "iso": iso,
            "epargne": clean(row[2]) if len(row) > 2 else None,
            "pret": clean(row[4]) if len(row) > 4 else None,
            "droit_adhesion": clean(row[5]) if len(row) > 5 else None,
            "parts_sociales": clean(row[6]) if len(row) > 6 else None,
            "total_fichier": total_fichier, "montant_vir": montant_vir,
            "ecart": round(ecart), "regle": abs(ecart) < 1, "observation": obs,
        })

    tp = sum(l["total_fichier"] or 0 for l in lignes)
    tr = sum(l["montant_vir"] or 0 for l in lignes)
    att = [l for l in lignes if not l["regle"]]
    return {
        "meta": {"titre": "Suivi des prélèvements & reversements SIVE", "exercice": 2026, "source": "POINT SIVE.xlsx"},
        "kpi": {"total_preleve": round(tp), "total_reverse": round(tr),
                "ecart_total": round(tp - tr), "nb_mois_attente": len(att),
                "taux_reversement": round(tr / tp * 100, 1) if tp else 0},
        "lignes": lignes,
    }


def main():
    data = parse_sive(SRC)
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    k = data["kpi"]
    print("REVERSEMENTS SIVE ANALYSÉS")
    print(f"  Total prélevé   : {k['total_preleve']:,} FCFA")
    print(f"  Total reversé   : {k['total_reverse']:,} FCFA")
    print(f"  ÉCART EN ATTENTE: {k['ecart_total']:,} FCFA ({k['nb_mois_attente']} mois)")
    print(f"  -> {os.path.abspath(OUT)}")


if __name__ == "__main__":
    main()
