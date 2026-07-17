"""
MA2E — Pilote Trésorerie
Parser robuste du fichier "POINT RETRAIT" (point de trésorerie quotidien).

Transforme un classeur Excel hétérogène (41 onglets, formats variables,
montants en texte, erreurs de saisie) en une série temporelle propre,
au format long : une ligne par (jour, compte).

Défis gérés :
  - Schéma variable selon l'année (colonne JULAYA apparue en 2026)
    -> mapping des colonnes par NOM d'en-tête, jamais par position.
  - Montants stockés en texte ("  1,421,423,504  ") ou en nombre.
  - Lignes parasites (totaux "-", "0", libellés SEMAINE, bas de tableau).
  - Erreurs de saisie : année incohérente, dates en texte ("14/05/206").

Sorties :
  - output/soldes_journaliers.csv   (données propres, format long)
  - output/rapport_qualite.txt      (rapport de contrôle qualité)

Usage : python parse_tresorerie.py "../../POINT RETRAIT AU 12-06_2026.xlsx"
"""
from __future__ import annotations
import sys, os, re, csv, datetime as dt
from collections import Counter

import openpyxl

# Console Windows (cp1252) : éviter les plantages d'encodage à l'affichage
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# --- Référentiel des comptes : motif d'en-tête -> (code, type) -------
# L'ordre n'a pas d'importance : on identifie par le libellé de l'en-tête.
COMPTES = [
    (r"ECOBANK",              "ECOBANK",          "BANQUE"),
    (r"ORABANK",              "ORABANK",          "BANQUE"),
    (r"BOA",                  "BOA-CI",           "BANQUE"),
    (r"UBA",                  "UBA",              "BANQUE"),
    (r"JULAYA",               "JULAYA",           "BANQUE"),
    (r"CAISSE.*DEPENSE",      "CAISSE_DEPENSES",  "CAISSE"),
    (r"CAISSE.*RECETTE",      "CAISSE_RECETTE",   "CAISSE"),
    (r"COFFRE",               "COFFRE_FORT",      "COFFRE"),
]


def clean_amount(value):
    """Convertit une cellule (texte ou nombre) en float, ou None si vide."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    # séparateurs : espaces (dont insécables), virgules de milliers
    s = s.replace(" ", "").replace(" ", "").replace(",", "")
    if s in ("", "-", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_date(value, expected_year):
    """Renvoie (date|None, suspecte: bool). Détecte les anomalies de saisie."""
    if isinstance(value, dt.datetime):
        d = value.date()
        # année incohérente avec l'onglet (ex : 2023-12-08 dans Janvier 2024)
        suspecte = (expected_year is not None and d.year != expected_year)
        return d, suspecte
    if isinstance(value, dt.date):
        suspecte = (expected_year is not None and value.year != expected_year)
        return value, suspecte
    if isinstance(value, str):
        s = value.strip()
        m = re.match(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
        if m:
            j, mo, an = (int(x) for x in m.groups())
            if an < 100:
                an += 2000
            # typo type "206" -> on recale sur l'année attendue
            if an < 1000 and expected_year:
                an = expected_year
            try:
                return dt.date(an, mo, j), True  # date en texte = toujours à vérifier
            except ValueError:
                return None, True
    return None, False


def year_from_title(cell, sheet_name):
    """Extrait l'année (R0 'TRESORERIE 2025' ou nom d'onglet)."""
    for src in (str(cell or ""), sheet_name):
        m = re.search(r"(20\d{2})", src)
        if m:
            return int(m.group(1))
    return None


def find_header(rows):
    """Trouve la ligne d'en-tête (celle qui contient 'ECOBANK')."""
    for i, row in enumerate(rows[:8]):
        joined = " ".join(str(c) for c in row if c is not None).upper()
        if "ECOBANK" in joined:
            return i
    return None


def map_columns(header_row):
    """Mappe index de colonne -> (code, type) + repère date et total."""
    bank_cols, date_col, total_col = {}, None, None
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        label = str(cell).upper()
        if re.search(r"JOUR|JOURN", label) and date_col is None:
            date_col = idx
        if "TOTAL" in label:
            total_col = idx
        for pattern, code, type_c in COMPTES:
            if re.search(pattern, label):
                bank_cols[idx] = (code, type_c)
                break
    return bank_cols, date_col, total_col


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    fichier = os.path.basename(path)
    lignes_propres = []           # dicts: date, code, type, montant, onglet, ligne, flag
    comptes_vus = {}              # code -> type
    stats = Counter()
    anomalies = []                # messages de contrôle qualité

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        annee = year_from_title(rows[0][3] if len(rows[0]) > 3 else None, sheet_name)
        hidx = find_header(rows)
        if hidx is None:
            stats["onglets_sans_entete"] += 1
            continue
        bank_cols, date_col, total_col = map_columns(rows[hidx])
        if not bank_cols or date_col is None:
            stats["onglets_non_mappes"] += 1
            continue
        stats["onglets_traites"] += 1

        for r in range(hidx + 1, len(rows)):
            row = rows[r]
            if date_col >= len(row):
                continue
            d, suspecte = parse_date(row[date_col], annee)
            if d is None:
                continue  # ligne sans date = séparateur / bas de tableau -> ignorée
            stats["jours_lus"] += 1
            flag_jour = "DATE_SUSPECTE" if suspecte else "OK"
            if suspecte:
                anomalies.append(
                    f"[{sheet_name}] L{r+1}: date suspecte {d} "
                    f"(année attendue {annee}) — valeur brute {row[date_col]!r}")

            # contrôle d'intégrité : somme des comptes vs colonne TOTAL
            somme = 0.0
            for idx, (code, type_c) in bank_cols.items():
                montant = clean_amount(row[idx]) if idx < len(row) else None
                if montant is None:
                    continue
                comptes_vus[code] = type_c
                somme += montant
                lignes_propres.append({
                    "date_jour": d.isoformat(),
                    "code_compte": code,
                    "type_compte": type_c,
                    "montant": round(montant, 2),
                    "source_onglet": sheet_name.strip(),
                    "source_ligne": r + 1,
                    "qualite_flag": flag_jour,
                })
                stats["soldes_extraits"] += 1

            if total_col is not None and total_col < len(row):
                total_excel = clean_amount(row[total_col])
                if total_excel and somme and abs(total_excel - somme) > 1:
                    stats["ecarts_total"] += 1
                    anomalies.append(
                        f"[{sheet_name}] L{r+1} ({d}): écart TOTAL — "
                        f"calculé {somme:,.0f} vs fichier {total_excel:,.0f} "
                        f"(diff {somme-total_excel:,.0f})")

    return lignes_propres, comptes_vus, stats, anomalies, fichier


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "../../POINT RETRAIT AU 12-06_2026.xlsx"
    here = os.path.dirname(os.path.abspath(__file__))
    out_dir = os.path.join(here, "..", "output")
    os.makedirs(out_dir, exist_ok=True)

    lignes, comptes, stats, anomalies, fichier = parse_workbook(path)

    # --- CSV propre (format long) ---
    csv_path = os.path.join(out_dir, "soldes_journaliers.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[
            "date_jour", "code_compte", "type_compte", "montant",
            "source_onglet", "source_ligne", "qualite_flag"])
        w.writeheader()
        w.writerows(lignes)

    # --- Rapport qualité ---
    dates = sorted({l["date_jour"] for l in lignes})
    rpt_path = os.path.join(out_dir, "rapport_qualite.txt")
    with open(rpt_path, "w", encoding="utf-8") as fh:
        fh.write("RAPPORT DE QUALITÉ — INGESTION POINT DE TRÉSORERIE\n")
        fh.write("=" * 60 + "\n")
        fh.write(f"Fichier source       : {fichier}\n")
        fh.write(f"Onglets traités      : {stats['onglets_traites']}\n")
        fh.write(f"Jours (lignes) lus   : {stats['jours_lus']}\n")
        fh.write(f"Soldes extraits      : {stats['soldes_extraits']}\n")
        fh.write(f"Comptes détectés     : {', '.join(sorted(comptes))}\n")
        if dates:
            fh.write(f"Période couverte     : {dates[0]} -> {dates[-1]} "
                     f"({len(dates)} jours distincts)\n")
        fh.write(f"Écarts TOTAL détectés: {stats['ecarts_total']}\n")
        fh.write(f"Anomalies totales    : {len(anomalies)}\n")
        fh.write("\n--- Détail des anomalies (50 premières) ---\n")
        for a in anomalies[:50]:
            fh.write(a + "\n")

    # --- Résumé console ---
    print("INGESTION TERMINÉE")
    print(f"  Onglets traités : {stats['onglets_traites']}")
    print(f"  Soldes extraits : {stats['soldes_extraits']:,}")
    print(f"  Comptes         : {', '.join(sorted(comptes))}")
    if dates:
        print(f"  Période         : {dates[0]} -> {dates[-1]} ({len(dates)} jours)")
    print(f"  Anomalies       : {len(anomalies)} (dont {stats['ecarts_total']} écarts de total)")
    print(f"  -> {csv_path}")
    print(f"  -> {rpt_path}")


if __name__ == "__main__":
    main()
